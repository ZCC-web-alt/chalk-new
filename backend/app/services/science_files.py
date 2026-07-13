from __future__ import annotations

import csv
import hashlib
import io
import math
import zipfile
from datetime import date, datetime, time
from itertools import islice
from pathlib import Path
from typing import Any

from PIL import Image, UnidentifiedImageError


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".gif", ".webp", ".tif", ".tiff"}
WORKBOOK_EXTENSIONS = {".xlsx", ".xls", ".csv"}
MAX_IMAGE_PIXELS = 50_000_000
MAX_XLSX_MEMBERS = 10_000
MAX_XLSX_UNCOMPRESSED_BYTES = 250 * 1024 * 1024
MAX_PREVIEW_COLUMNS = 200
MAX_PREVIEW_CELL_CHARS = 2000


class InvalidScienceFile(ValueError):
    pass


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def inspect_multimodal_file(path: Path, original_name: str) -> tuple[str, str, dict[str, Any]]:
    suffix = Path(original_name).suffix.lower()
    if suffix in IMAGE_EXTENSIONS:
        return _inspect_image(path, suffix)
    if suffix == ".xlsx":
        return _inspect_xlsx(path)
    if suffix == ".xls":
        return _inspect_xls(path)
    if suffix == ".csv":
        return _inspect_csv(path)
    raise InvalidScienceFile("Unsupported multimodal file type.")


def _inspect_image(path: Path, suffix: str) -> tuple[str, str, dict[str, Any]]:
    try:
        with Image.open(path) as image:
            width, height = image.size
            if width < 1 or height < 1 or width * height > MAX_IMAGE_PIXELS:
                raise InvalidScienceFile("The image dimensions are not allowed.")
            image.verify()
        with Image.open(path) as image:
            image.load()
            image_format = str(image.format or "").upper()
            frame_count = int(getattr(image, "n_frames", 1) or 1)
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
        raise InvalidScienceFile("The uploaded image could not be decoded.") from exc

    expected = {
        ".jpg": "JPEG", ".jpeg": "JPEG", ".png": "PNG", ".bmp": "BMP",
        ".gif": "GIF", ".webp": "WEBP", ".tif": "TIFF", ".tiff": "TIFF",
    }[suffix]
    if image_format != expected:
        raise InvalidScienceFile("The image content does not match its extension.")
    mime = Image.MIME.get(image_format, "application/octet-stream")
    warnings = ["Only the first GIF frame will be analyzed."] if image_format == "GIF" and frame_count > 1 else []
    return "image", mime, {
        "width": width,
        "height": height,
        "format": image_format,
        "frameCount": frame_count,
        "warnings": warnings,
    }


def _inspect_xlsx(path: Path) -> tuple[str, str, dict[str, Any]]:
    with path.open("rb") as handle:
        header = handle.read(4)
    if header != b"PK\x03\x04":
        raise InvalidScienceFile("The workbook content does not match XLSX format.")
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_MEMBERS:
                raise InvalidScienceFile("The workbook contains too many archive members.")
            if sum(member.file_size for member in members) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise InvalidScienceFile("The workbook expands beyond the configured limit.")
            if not any(member.filename == "xl/workbook.xml" for member in members):
                raise InvalidScienceFile("The XLSX workbook manifest is missing.")
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False)
        sheet_names = list(workbook.sheetnames)
        sheet_rows = {sheet.title: max(int(sheet.max_row or 0) - 1, 0) for sheet in workbook.worksheets}
        sheet_columns = {sheet.title: int(sheet.max_column or 0) for sheet in workbook.worksheets}
        workbook.close()
    except (OSError, ValueError, zipfile.BadZipFile, KeyError) as exc:
        if isinstance(exc, InvalidScienceFile):
            raise
        raise InvalidScienceFile("The XLSX workbook could not be read.") from exc
    return "workbook", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", {
        "sheetNames": sheet_names,
        "sheetCount": len(sheet_names),
        "sheetRows": sheet_rows,
        "sheetColumns": sheet_columns,
    }


def _inspect_xls(path: Path) -> tuple[str, str, dict[str, Any]]:
    with path.open("rb") as handle:
        header = handle.read(8)
    if header != bytes.fromhex("D0CF11E0A1B11AE1"):
        raise InvalidScienceFile("The workbook content does not match XLS format.")
    try:
        import xlrd

        workbook = xlrd.open_workbook(path, on_demand=True)
        sheet_names = list(workbook.sheet_names())
        sheet_rows = {name: max(int(workbook.sheet_by_name(name).nrows) - 1, 0) for name in sheet_names}
        sheet_columns = {name: int(workbook.sheet_by_name(name).ncols) for name in sheet_names}
        workbook.release_resources()
    except (ImportError, OSError, ValueError) as exc:
        raise InvalidScienceFile("The XLS workbook could not be read.") from exc
    return "workbook", "application/vnd.ms-excel", {
        "sheetNames": sheet_names,
        "sheetCount": len(sheet_names),
        "sheetRows": sheet_rows,
        "sheetColumns": sheet_columns,
    }


def _inspect_csv(path: Path) -> tuple[str, str, dict[str, Any]]:
    raw = path.read_bytes()
    decoded = None
    encoding = ""
    for candidate in ("utf-8-sig", "gb18030"):
        try:
            decoded = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if decoded is None or "\x00" in decoded:
        raise InvalidScienceFile("The CSV file encoding is not supported.")
    row_count = 0
    max_columns = 0
    for row in csv.reader(io.StringIO(decoded)):
        row_count += 1
        max_columns = max(max_columns, len(row))
    return "workbook", "text/csv", {
        "sheetNames": ["CSV"],
        "sheetCount": 1,
        "rows": row_count,
        "dataRows": max(row_count - 1, 0),
        "columns": max_columns,
        "encoding": encoding,
    }


def read_table_preview(
    path: Path,
    *,
    metadata: dict[str, Any],
    sheet_name: str | None,
    page: int,
    page_size: int,
) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        selected_sheet = "CSV"
        if sheet_name and sheet_name != selected_sheet:
            raise InvalidScienceFile("The selected CSV sheet does not exist.")
        columns, rows, total_rows = _read_csv_page(
            path,
            encoding=str(metadata.get("encoding") or "utf-8-sig"),
            page=page,
            page_size=page_size,
        )
    elif suffix == ".xlsx":
        columns, rows, total_rows, selected_sheet = _read_xlsx_page(
            path,
            sheet_name=sheet_name,
            page=page,
            page_size=page_size,
        )
    elif suffix == ".xls":
        columns, rows, total_rows, selected_sheet = _read_xls_page(
            path,
            sheet_name=sheet_name,
            page=page,
            page_size=page_size,
        )
    else:
        raise InvalidScienceFile("This asset does not support table preview.")
    total_pages = math.ceil(total_rows / page_size) if total_rows else 0
    sheet_columns = metadata.get("sheetColumns") if isinstance(metadata.get("sheetColumns"), dict) else {}
    source_column_count = int(sheet_columns.get(selected_sheet) or metadata.get("columns") or len(columns))
    return {
        "sheetName": selected_sheet,
        "columns": columns,
        "rows": rows,
        "pagination": {
            "page": page,
            "pageSize": page_size,
            "totalItems": total_rows,
            "totalPages": total_pages,
        },
        "coverage": {
            "rowsRead": total_rows,
            "columnsRead": len(columns),
            "columnsTruncated": source_column_count > len(columns),
            "strategy": "paged-read-only-preview",
        },
    }


def _read_csv_page(path: Path, *, encoding: str, page: int, page_size: int):
    start = (page - 1) * page_size
    selected: list[list[Any]] = []
    total_rows = 0
    try:
        with path.open("r", encoding=encoding, newline="") as handle:
            reader = csv.reader(handle)
            header = next(reader, [])
            columns = _preview_columns(header)
            for index, row in enumerate(reader):
                if start <= index < start + page_size:
                    selected.append(_preview_row(row, len(columns)))
                total_rows += 1
    except (OSError, UnicodeError, csv.Error) as exc:
        raise InvalidScienceFile("The CSV preview could not be read.") from exc
    return columns, selected, total_rows


def _read_xlsx_page(path: Path, *, sheet_name: str | None, page: int, page_size: int):
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        selected_sheet = sheet_name or (workbook.sheetnames[0] if workbook.sheetnames else "")
        if selected_sheet not in workbook.sheetnames:
            workbook.close()
            raise InvalidScienceFile("The selected workbook sheet does not exist.")
        worksheet = workbook[selected_sheet]
        iterator = worksheet.iter_rows(values_only=True)
        columns = _preview_columns(next(iterator, ()))
        start = (page - 1) * page_size
        rows = [_preview_row(row, len(columns)) for row in islice(iterator, start, start + page_size)]
        total_rows = max(int(worksheet.max_row or 0) - 1, 0)
        workbook.close()
        return columns, rows, total_rows, selected_sheet
    except InvalidScienceFile:
        raise
    except (OSError, ValueError, KeyError) as exc:
        raise InvalidScienceFile("The XLSX preview could not be read.") from exc


def _read_xls_page(path: Path, *, sheet_name: str | None, page: int, page_size: int):
    try:
        import xlrd

        workbook = xlrd.open_workbook(path, on_demand=True)
        selected_sheet = sheet_name or (workbook.sheet_names()[0] if workbook.sheet_names() else "")
        if selected_sheet not in workbook.sheet_names():
            workbook.release_resources()
            raise InvalidScienceFile("The selected workbook sheet does not exist.")
        worksheet = workbook.sheet_by_name(selected_sheet)
        columns = _preview_columns(worksheet.row_values(0) if worksheet.nrows else [])
        start = (page - 1) * page_size + 1
        stop = min(start + page_size, worksheet.nrows)
        rows = [_preview_row(worksheet.row_values(index), len(columns)) for index in range(start, stop)]
        total_rows = max(int(worksheet.nrows) - 1, 0)
        workbook.release_resources()
        return columns, rows, total_rows, selected_sheet
    except InvalidScienceFile:
        raise
    except (ImportError, OSError, ValueError, IndexError) as exc:
        raise InvalidScienceFile("The XLS preview could not be read.") from exc


def _preview_columns(values) -> list[str]:
    columns = []
    for index, value in enumerate(list(values)[:MAX_PREVIEW_COLUMNS], 1):
        text = str(value or "").strip()
        columns.append(text[:240] or f"Column {index}")
    return columns


def _preview_row(values, column_count: int) -> list[Any]:
    row = [_preview_cell(value) for value in list(values)[:column_count]]
    return row + [None] * max(0, column_count - len(row))


def _preview_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if hasattr(value, "item"):
        try:
            return _preview_cell(value.item())
        except (TypeError, ValueError):
            pass
    text = str(value)
    return text[:MAX_PREVIEW_CELL_CHARS]
