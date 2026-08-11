from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import tempfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import UTC, datetime
from functools import lru_cache
from pathlib import Path
from threading import RLock
from typing import Any, Iterable, Mapping, Sequence

from app.core.config import get_settings
from app.core.legacy import llm_client as load_llm_client
from app.services import api_keys
from app.services.model_call_ledger import ModelCallLedgerStore
from app.services.research_generation import (
    ResearchGenerationCallError,
    ResearchGenerationRequest,
    ResearchGenerationService,
    ResearchGenerationValidationError,
)
from app.services.science125_catalog import (
    get_science125_catalog,
    get_science125_route,
    get_science125_routing,
)
from app.services.science125_context import load_science125_context_index
from app.services.science125_localization import get_science125_localization
from app.services.science125_prompts import (
    SCIENCE125_PROMPT_VERSION,
    load_science125_prompt_registry,
)
from app.services.science125_queries import build_science125_refinement_queries
from app.services.science125_relevance import EvidenceQualification, qualify_science125_evidence
from app.services.science125_report_store import (
    Science125ReportStore,
    StoredScience125Batch,
    StoredScience125Export,
    StoredScience125Report,
)
from app.services.job_store import StoredJob
from app.schemas.research import ResearchOutput
from sqlalchemy.exc import IntegrityError
from app.services.science125_retrieval import (
    EvidenceRecord,
    ProviderRateStateStore,
    default_provider_adapters,
    get_provider,
    get_science125_retrieval_profile,
    search_science125,
)


SCIENCE125_REPORT_EXPORT_DIR = "science125-report-exports"
_STOP_WORDS = {
    "about", "after", "again", "against", "among", "because", "before", "being", "between",
    "could", "does", "from", "have", "into", "itself", "level", "might", "other", "should",
    "their", "there", "these", "those", "through", "under", "using", "what", "when", "where",
    "which", "while", "with", "would", "your", "ours", "they", "them", "this", "that", "than",
    "then", "will", "ever", "make", "made", "much", "many", "more", "most", "some", "such",
    "only", "also", "very", "still", "were", "been", "being", "how", "can", "we", "the", "and",
    "for", "are", "but", "not", "you", "its", "our", "has", "had", "was", "why", "who",
}


class Science125ReportError(RuntimeError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def _canonical_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def _sha256_json(value: Any) -> str:
    return hashlib.sha256(_canonical_json(value).encode("utf-8")).hexdigest()


def _provider_family(provider_id: str) -> str:
    try:
        return get_provider(provider_id).family
    except KeyError:
        return str(provider_id or "unknown").strip().lower() or "unknown"


def _work_key(record: EvidenceRecord) -> str:
    for prefix, value in (
        ("doi", record.doi),
        ("pmid", record.pmid),
        ("arxiv", record.arxiv_id),
        ("ads", record.ads_id),
    ):
        if value:
            return f"{prefix}:{str(value).strip().casefold()}"
    return f"stable:{record.stable_id.strip().casefold()}"


@dataclass(frozen=True, slots=True)
class Science125SelectedHypothesis:
    hypothesis_id: str | None
    confidence: float
    reason: str

    def to_dict(self) -> dict[str, Any]:
        return {
            "hypothesisId": self.hypothesis_id,
            "confidence": self.confidence,
            "reason": self.reason,
        }


@dataclass(frozen=True, slots=True)
class Science125EvidenceSelectionItem:
    stable_id: str
    provider: str
    provider_family: str
    title: str
    record: dict[str, Any]
    relevance_score: float
    relevance_label: str
    relevance_breakdown: dict[str, Any]
    eligible: bool
    selected: bool
    reasons: tuple[str, ...]

    def to_dict(self) -> dict[str, Any]:
        return {
            "stableId": self.stable_id,
            "provider": self.provider,
            "providerFamily": self.provider_family,
            "title": self.title,
            "record": self.record,
            "relevanceScore": self.relevance_score,
            "relevanceLabel": self.relevance_label,
            "relevanceBreakdown": self.relevance_breakdown,
            "eligibleForGeneration": self.eligible,
            "selected": self.selected,
            "reasons": list(self.reasons),
        }


@dataclass(frozen=True, slots=True)
class Science125EvidenceSelection:
    evidence_status: str
    selected: tuple[Science125EvidenceSelectionItem, ...]
    excluded: tuple[Science125EvidenceSelectionItem, ...]
    eligible_full_text_count: int
    provider_family_count: int
    provider_families: tuple[str, ...]
    evidence_snapshot_hash: str

    def evidence_records(self) -> tuple[dict[str, Any], ...]:
        return tuple(
            {
                **item.record,
                "stableId": item.stable_id,
                "providerFamily": item.provider_family,
                "selectionReason": item.reasons[0] if item.reasons else "AUTO_SELECTED",
                "relevanceScore": item.relevance_score,
                "relevanceLabel": item.relevance_label,
            }
            for item in self.selected
        )

    def to_snapshot(self) -> dict[str, Any]:
        return {
            "selectionVersion": "science125-auto-evidence-selection-v1",
            "evidenceStatus": self.evidence_status,
            "eligibleFullTextCount": self.eligible_full_text_count,
            "providerFamilyCount": self.provider_family_count,
            "providerFamilies": list(self.provider_families),
            "selectedEvidence": [item.to_dict() for item in self.selected],
            "excludedEvidence": [item.to_dict() for item in self.excluded],
        }


def select_science125_hypothesis(research_output: Mapping[str, Any]) -> Science125SelectedHypothesis:
    raw_hypotheses = research_output.get("hypotheses")
    hypotheses = raw_hypotheses if isinstance(raw_hypotheses, Sequence) and not isinstance(raw_hypotheses, (str, bytes)) else []
    candidates: list[dict[str, Any]] = []
    for raw in hypotheses:
        if not isinstance(raw, Mapping):
            continue
        hypothesis_id = str(raw.get("id") or "").strip()
        if not re.fullmatch(r"H[1-5]", hypothesis_id):
            continue
        try:
            confidence = float(raw.get("confidence") or 0.0)
        except (TypeError, ValueError):
            confidence = 0.0
        if confidence <= 0:
            continue
        supporting = raw.get("supportingEvidenceRefs")
        falsification = raw.get("falsificationCriteria")
        candidates.append({
            "id": hypothesis_id,
            "confidence": confidence,
            "supporting": len(supporting) if isinstance(supporting, Sequence) and not isinstance(supporting, (str, bytes)) else 0,
            "falsification": len(falsification) if isinstance(falsification, Sequence) and not isinstance(falsification, (str, bytes)) else 0,
        })
    if not candidates:
        return Science125SelectedHypothesis(
            hypothesis_id=None,
            confidence=0.0,
            reason="没有候选假设的置信度大于 0；H0 仅作为零假设对照，不参与自动选中。",
        )
    selected = sorted(
        candidates,
        key=lambda item: (
            -float(item["confidence"]),
            -int(item["supporting"]),
            -int(item["falsification"]),
            int(str(item["id"])[1:]),
        ),
    )[0]
    return Science125SelectedHypothesis(
        hypothesis_id=str(selected["id"]),
        confidence=float(selected["confidence"]),
        reason=(
            f"系统按置信度最高、证据引用更多、可证伪条件更完整、ID 更小的顺序自动选择 {selected['id']}；"
            "H0 保留为零假设对照，不参与选择。"
        ),
    )


def _selection_item(
    *,
    record: EvidenceRecord,
    qualification: EvidenceQualification,
    provider_family: str,
    selected: bool,
    reasons: Iterable[str],
) -> Science125EvidenceSelectionItem:
    return Science125EvidenceSelectionItem(
        stable_id=record.stable_id,
        provider=record.provider,
        provider_family=provider_family,
        title=record.title,
        record=record.to_dict(),
        relevance_score=qualification.relevance.score,
        relevance_label=qualification.relevance.label,
        relevance_breakdown=qualification.relevance.to_dict(),
        eligible=qualification.eligible_for_generation,
        selected=selected,
        reasons=tuple(str(reason) for reason in reasons),
    )


def select_science125_evidence(
    *,
    question_id: str,
    query: str,
    records: Sequence[EvidenceRecord],
    max_selected: int = 6,
) -> Science125EvidenceSelection:
    best_by_work: dict[str, tuple[EvidenceRecord, EvidenceQualification, str]] = {}
    duplicate_items: list[Science125EvidenceSelectionItem] = []
    excluded_items: list[Science125EvidenceSelectionItem] = []
    for record in records:
        qualification = qualify_science125_evidence(question_id, query, record)
        family = _provider_family(record.provider)
        key = _work_key(record)
        existing = best_by_work.get(key)
        if existing is None or qualification.relevance.score > existing[1].relevance.score:
            if existing is not None:
                duplicate_items.append(_selection_item(
                    record=existing[0],
                    qualification=existing[1],
                    provider_family=existing[2],
                    selected=False,
                    reasons=("DUPLICATE_WORK",),
                ))
            best_by_work[key] = (record, qualification, family)
        else:
            duplicate_items.append(_selection_item(
                record=record,
                qualification=qualification,
                provider_family=family,
                selected=False,
                reasons=("DUPLICATE_WORK",),
            ))
    eligible: list[tuple[EvidenceRecord, EvidenceQualification, str]] = []
    for record, qualification, family in best_by_work.values():
        if qualification.eligible_for_generation:
            eligible.append((record, qualification, family))
        else:
            excluded_items.append(_selection_item(
                record=record,
                qualification=qualification,
                provider_family=family,
                selected=False,
                reasons=qualification.reasons,
            ))
    ranked = sorted(
        eligible,
        key=lambda item: (
            -item[1].relevance.score,
            item[0].provider,
            item[0].stable_id,
        ),
    )
    eligible_families = sorted({family for _record, _qualification, family in eligible})
    if len(eligible) < 3 or len(eligible_families) < 2:
        selected_items: list[Science125EvidenceSelectionItem] = []
        excluded_items.extend(
            _selection_item(
                record=record,
                qualification=qualification,
                provider_family=family,
                selected=False,
                reasons=("GATE_FAILED_NEEDS_3_FULL_TEXT_AND_2_FAMILIES",),
            )
            for record, qualification, family in ranked
        )
        evidence_status = "insufficient"
    else:
        chosen: list[tuple[EvidenceRecord, EvidenceQualification, str]] = []
        seen_keys: set[str] = set()
        seen_families: set[str] = set()
        for record, qualification, family in ranked:
            if family in seen_families:
                continue
            chosen.append((record, qualification, family))
            seen_keys.add(_work_key(record))
            seen_families.add(family)
            if len(chosen) >= min(max_selected, len(ranked)):
                break
        for record, qualification, family in ranked:
            if len(chosen) >= max_selected:
                break
            key = _work_key(record)
            if key in seen_keys:
                continue
            chosen.append((record, qualification, family))
            seen_keys.add(key)
        selected_items = [
            _selection_item(
                record=record,
                qualification=qualification,
                provider_family=family,
                selected=True,
                reasons=(f"AUTO_SELECTED_MEDIUM_OR_HIGH_FULL_TEXT relevance={qualification.relevance.score:.4f} family={family}",),
            )
            for record, qualification, family in chosen
        ]
        chosen_keys = {_work_key(record) for record, _qualification, _family in chosen}
        excluded_items.extend(
            _selection_item(
                record=record,
                qualification=qualification,
                provider_family=family,
                selected=False,
                reasons=("NOT_IN_TOP_AUTO_SELECTED_SET",),
            )
            for record, qualification, family in ranked
            if _work_key(record) not in chosen_keys
        )
        evidence_status = "sufficient"
    excluded_items.extend(duplicate_items)
    snapshot_payload = {
        "questionId": question_id,
        "query": query,
        "selectedStableIds": [item.stable_id for item in selected_items],
        "excludedStableIds": [item.stable_id for item in excluded_items],
        "eligibleFullTextCount": len(eligible),
        "providerFamilies": eligible_families,
    }
    return Science125EvidenceSelection(
        evidence_status=evidence_status,
        selected=tuple(selected_items),
        excluded=tuple(excluded_items),
        eligible_full_text_count=len(eligible),
        provider_family_count=len(eligible_families),
        provider_families=tuple(eligible_families),
        evidence_snapshot_hash=_sha256_json(snapshot_payload),
    )


def science125_exports_dir() -> Path:
    root = get_settings().data_dir / SCIENCE125_REPORT_EXPORT_DIR
    root.mkdir(parents=True, exist_ok=True)
    return root


def science125_model_name() -> str:
    return str(getattr(load_llm_client(), "REASONING_MODEL", "qwen3.8-max") or "qwen3.8-max")


def build_science125_search_query(question: str, source_context: str) -> str:
    tokens = re.findall(r"[a-z][a-z0-9]*(?:[-'][a-z0-9]+)*", f"{question}\n{source_context}".casefold())
    selected: list[str] = []
    seen: set[str] = set()
    for token in tokens:
        if len(token) < 4 or token in _STOP_WORDS or token in seen:
            continue
        seen.add(token)
        selected.append(token)
        if len(selected) >= 16:
            break
    return " ".join(selected) if len(selected) >= 4 else question.strip()


def _report_payload(report: StoredScience125Report, exports: tuple[dict[str, Any], ...] = ()) -> dict[str, Any]:
    return report.to_dict(exports=exports)


def _batch_payload(batch: StoredScience125Batch, reports: tuple[StoredScience125Report, ...]) -> dict[str, Any]:
    return batch.to_dict(reports=tuple(report.to_summary_dict() for report in reports))


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    temporary.replace(path)


def _text(value: Any) -> str:
    return str(value or "").strip()


def _hypothesis_title(hypothesis: Mapping[str, Any]) -> str:
    return _text(hypothesis.get("title")) or _text(hypothesis.get("statement"))[:80] or _text(hypothesis.get("id"))


def _add_key_value(doc: Any, key: str, value: Any) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run(f"{key}: ").bold = True
    paragraph.add_run(_text(value) or "未记录")


def _build_report_docx(path: Path, report: StoredScience125Report) -> None:
    from docx import Document

    output = report.research_output or {}
    science125 = output.get("science125") if isinstance(output.get("science125"), Mapping) else {}
    doc = Document()
    doc.add_heading(f"Science 125 假设报告：{report.question_id}", level=0)
    _add_key_value(doc, "英文原题", report.question)
    _add_key_value(doc, "中文译题", report.question_zh or "")
    _add_key_value(doc, "领域", f"{report.benchmark_domain} / {report.primary_subdomain}")
    _add_key_value(doc, "方法类型", json.dumps(science125.get("methodProfile") or {}, ensure_ascii=False))
    _add_key_value(doc, "自动选中假设", report.selected_hypothesis_id or "无")
    _add_key_value(doc, "选择原因", report.selected_hypothesis_reason or "无")

    doc.add_heading("自动选择的文献证据", level=1)
    selected = (report.evidence_snapshot or {}).get("selectedEvidence") or []
    if selected:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for cell, title in zip(table.rows[0].cells, ("Provider", "Title", "ID", "Full text", "Reason"), strict=True):
            cell.text = title
        for item in selected:
            record = item.get("record") if isinstance(item, Mapping) else {}
            row = table.add_row().cells
            row[0].text = _text(item.get("provider") if isinstance(item, Mapping) else "")
            row[1].text = _text(item.get("title") if isinstance(item, Mapping) else "")
            row[2].text = _text(item.get("stableId") if isinstance(item, Mapping) else "")
            row[3].text = _text(record.get("accessStatus") if isinstance(record, Mapping) else "")
            row[4].text = "; ".join(str(reason) for reason in item.get("reasons", [])) if isinstance(item, Mapping) else ""
    else:
        doc.add_paragraph("未达到证据门禁，未生成正式文献证据表。")

    brief = output.get("brief") if isinstance(output.get("brief"), Mapping) else {}
    doc.add_heading("研究简报", level=1)
    _add_key_value(doc, "研究问题", brief.get("researchQuestion"))
    _add_key_value(doc, "背景", brief.get("background"))
    _add_key_value(doc, "范围", brief.get("scope"))

    doc.add_heading("候选假设", level=1)
    for hypothesis in output.get("hypotheses") or []:
        if not isinstance(hypothesis, Mapping):
            continue
        label = f"{hypothesis.get('id')} - {_hypothesis_title(hypothesis)}"
        if hypothesis.get("id") == report.selected_hypothesis_id:
            label += "（系统自动选中）"
        doc.add_heading(label, level=2)
        _add_key_value(doc, "陈述", hypothesis.get("statement"))
        _add_key_value(doc, "机制", hypothesis.get("mechanism"))
        _add_key_value(doc, "置信度", hypothesis.get("confidence"))
        _add_key_value(doc, "可证伪条件", "；".join(str(item) for item in hypothesis.get("falsificationCriteria", [])))

    null_hypothesis = output.get("nullHypothesis") if isinstance(output.get("nullHypothesis"), Mapping) else {}
    doc.add_heading("H0 零假设", level=1)
    _add_key_value(doc, "陈述", null_hypothesis.get("statement"))
    _add_key_value(doc, "机制", null_hypothesis.get("mechanism"))

    plan = output.get("researchPlan") if isinstance(output.get("researchPlan"), Mapping) else {}
    doc.add_heading("研究计划", level=1)
    for key, label in (
        ("independentVariables", "自变量"),
        ("dependentVariables", "因变量"),
        ("controlVariables", "控制变量"),
        ("decisionThresholds", "判定阈值"),
        ("stopConditions", "停止条件"),
        ("risks", "风险"),
        ("uncertainties", "不确定性"),
        ("applicabilityBoundaries", "适用边界"),
    ):
        _add_key_value(doc, label, "；".join(str(item) for item in plan.get(key, [])))

    doc.add_heading("模型与审计", level=1)
    _add_key_value(doc, "模型", report.model)
    _add_key_value(doc, "requestId", report.request_id)
    _add_key_value(doc, "证据快照哈希", report.evidence_snapshot_sha256)
    _add_key_value(doc, "生成时间", datetime.now(UTC).isoformat())
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _build_batch_docx(path: Path, batch: StoredScience125Batch, reports: Sequence[StoredScience125Report]) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("Science 125 全量假设汇总报告", level=0)
    _add_key_value(doc, "批次 ID", batch.id)
    _add_key_value(doc, "模型", batch.model)
    _add_key_value(doc, "完成情况", f"{batch.succeeded_count}/{batch.total_count}")
    _add_key_value(doc, "Manifest", f"{batch.manifest_version} / {batch.manifest_sha256}")
    _add_key_value(doc, "Routing", f"{batch.routing_version} / {batch.routing_sha256}")
    for report in sorted(reports, key=lambda item: item.question_id):
        doc.add_page_break()
        doc.add_heading(f"{report.question_id} {report.question_zh or report.question}", level=1)
        _add_key_value(doc, "英文原题", report.question)
        _add_key_value(doc, "领域", f"{report.benchmark_domain} / {report.primary_subdomain}")
        _add_key_value(doc, "自动选中假设", report.selected_hypothesis_id or "无")
        _add_key_value(doc, "选择原因", report.selected_hypothesis_reason or "无")
        output = report.research_output or {}
        hypotheses = output.get("hypotheses") if isinstance(output.get("hypotheses"), Sequence) else []
        for hypothesis in hypotheses:
            if isinstance(hypothesis, Mapping) and hypothesis.get("id") == report.selected_hypothesis_id:
                _add_key_value(doc, "假设陈述", hypothesis.get("statement"))
                _add_key_value(doc, "可证伪条件", "；".join(str(item) for item in hypothesis.get("falsificationCriteria", [])))
                break
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


class Science125ReportService:
    def __init__(self, *, store: Science125ReportStore | None = None):
        self._lock = RLock()
        self._store = store or Science125ReportStore(get_settings().web_db_path)
        self._rate_store = ProviderRateStateStore(self._store.path)
        self._ledger = ModelCallLedgerStore(self._store.path)
        self._executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="chalk-science125-batch")
        self._active_batch_ids: set[str] = set()

    @property
    def store(self) -> Science125ReportStore:
        return self._store

    def dispose(self) -> None:
        self._executor.shutdown(wait=False, cancel_futures=True)
        self._ledger.dispose()
        self._rate_store.dispose()
        self._store.dispose()

    def create_batch(self, *, user_id: int, question_ids: tuple[str, ...] | None = None) -> StoredScience125Batch:
        routing = get_science125_routing()
        prompt_registry = load_science125_prompt_registry(SCIENCE125_PROMPT_VERSION)
        return self._store.create_batch(
            user_id=user_id,
            question_ids=question_ids,
            manifest_version="science125-v1",
            manifest_sha256=routing.base_manifest_content_sha256,
            routing_version="science125-routing-v1",
            routing_sha256=routing.routing_content_sha256,
            prompt_version=SCIENCE125_PROMPT_VERSION,
            prompt_registry_sha256=prompt_registry.registry_content_sha256,
            model=science125_model_name(),
        )

    def list_batches(self, *, user_id: int) -> list[StoredScience125Batch]:
        return self._store.list_batches(user_id=user_id)

    def delete_batch(self, *, user_id: int, batch_id: str) -> StoredScience125Batch | None:
        with self._lock:
            batch = self._store.get_batch_for_user(user_id=user_id, batch_id=batch_id)
            if batch is None:
                return None
            if batch.status == "RUNNING" or batch_id in self._active_batch_ids:
                raise Science125ReportError(
                    "SCIENCE125_BATCH_DELETE_CONFLICT",
                    "A running Science 125 batch cannot be deleted. Pause it and wait for the current item to stop.",
                )
            deleted = self._store.delete_batch(user_id=user_id, batch_id=batch_id)
            if deleted is None:
                return None
            user_root = (science125_exports_dir() / str(user_id)).resolve()
            batch_root = (user_root / batch_id).resolve()
            if batch_root.is_relative_to(user_root):
                shutil.rmtree(batch_root, ignore_errors=True)
            return deleted

    def get_batch(self, *, user_id: int, batch_id: str) -> StoredScience125Batch | None:
        return self._store.get_batch_for_user(user_id=user_id, batch_id=batch_id)

    def batch_out(self, *, user_id: int, batch_id: str) -> dict[str, Any] | None:
        batch = self.get_batch(user_id=user_id, batch_id=batch_id)
        if batch is None:
            return None
        reports = tuple(self._store.list_reports(user_id=user_id, batch_id=batch_id, sort_by="questionId", sort_order="asc"))
        return _batch_payload(batch, reports)

    def list_reports(
        self,
        *,
        user_id: int,
        batch_id: str | None = None,
        question_id: str | None = None,
        status: str | None = None,
        benchmark_domain: str | None = None,
        sort_by: str = "updatedAt",
        sort_order: str = "desc",
    ) -> list[StoredScience125Report]:
        return self._store.list_reports(
            user_id=user_id,
            batch_id=batch_id,
            question_id=question_id,
            status=status,
            benchmark_domain=benchmark_domain,
            sort_by=sort_by,
            sort_order=sort_order,
        )

    def get_report(self, *, user_id: int, report_id: str) -> StoredScience125Report | None:
        return self._store.get_report_for_user(user_id=user_id, report_id=report_id)

    def import_interactive_job(
        self,
        *,
        job: StoredJob,
        literature_job: StoredJob | None = None,
    ) -> StoredScience125Report:
        """Adopt a completed interactive Science 125 job without retrieval or model calls."""
        if job.type != "hypothesis_generate" or job.status != "SUCCEEDED":
            raise Science125ReportError("SCIENCE125_JOB_NOT_IMPORTABLE", "Only successful hypothesis jobs can be imported.")
        question_id = str(job.payload.get("science125Id") or "").strip()
        if not question_id:
            raise Science125ReportError("SCIENCE125_ID_REQUIRED", "The completed job is not a Science 125 job.")
        existing = self._store.get_report_by_source_job(user_id=job.user_id, source_job_id=job.id)
        if existing is not None:
            self._ensure_report_exports(user_id=job.user_id, report=existing)
            return existing
        result = job.result if isinstance(job.result, dict) else {}
        private_snapshot = result.get("_reportSnapshot") if isinstance(result.get("_reportSnapshot"), dict) else {}
        output_data = result.get("researchOutput")
        audit = result.get("audit") if isinstance(result.get("audit"), dict) else {}
        try:
            output = ResearchOutput.model_validate(output_data)
        except Exception as exc:
            raise Science125ReportError("SCIENCE125_SCHEMA_INVALID", "The completed job does not contain valid research-v1 output.") from exc
        science125 = output.science125
        if science125 is None or science125.question_id != question_id:
            raise Science125ReportError("SCIENCE125_OUTPUT_MISMATCH", "The completed output belongs to another Science 125 item.")
        provider = str(audit.get("provider") or output.provenance.provider).strip()
        model = str(audit.get("model") or output.provenance.model).strip()
        if provider.casefold() != "dashscope" or not model.casefold().startswith("qwen"):
            raise Science125ReportError("SCIENCE125_QWEN_REQUIRED", "Science 125 reports must originate from DashScope Qwen.")
        if literature_job is None:
            raise Science125ReportError(
                "SCIENCE125_SEARCH_NOT_FOUND",
                "The completed job's literature snapshot is unavailable.",
            )
        if literature_job.user_id != job.user_id or literature_job.type != "literature_search" or literature_job.status != "SUCCEEDED":
            raise Science125ReportError("SCIENCE125_SEARCH_NOT_FOUND", "The completed literature search snapshot is unavailable.")
        if str(literature_job.payload.get("science125Id") or "") != question_id:
            raise Science125ReportError("SCIENCE125_SEARCH_MISMATCH", "The literature snapshot belongs to another item.")
        catalog = get_science125_catalog()
        item = next((entry for entry in catalog.data if entry.id == question_id), None)
        if item is None:
            raise Science125ReportError("SCIENCE125_QUESTION_NOT_FOUND", "Science 125 question not found.")
        route = get_science125_route(question_id)
        search_result = literature_job.result if isinstance(literature_job.result, dict) else {}
        query_info = search_result.get("query") if isinstance(search_result.get("query"), dict) else {}
        query = str(query_info.get("queryText") or literature_job.payload.get("queryText") or "").strip()
        rows = search_result.get("results") if isinstance(search_result.get("results"), list) else []
        reviewed_ids = [str(value).strip() for value in job.payload.get("reviewedEvidenceIds") or []]
        by_id = {str(row.get("id") or "").strip(): row for row in rows if isinstance(row, dict)}
        if not reviewed_ids or any(value not in by_id for value in reviewed_ids):
            raise Science125ReportError("SCIENCE125_EVIDENCE_SNAPSHOT_INVALID", "The reviewed evidence is missing from the completed search snapshot.")
        reviewed_snapshot_records = private_snapshot.get("reviewedEvidence") if isinstance(private_snapshot.get("reviewedEvidence"), list) else []
        evidence_records: list[EvidenceRecord] = []
        if reviewed_snapshot_records:
            for record in reviewed_snapshot_records:
                if not isinstance(record, dict):
                    continue
                evidence_records.append(EvidenceRecord(
                    provider=str(record.get("provider") or "").strip(),
                    stable_id=str(record.get("stableId") or "").strip(),
                    title=str(record.get("title") or "").strip(),
                    abstract=str(record.get("abstract") or "").strip(),
                    doi=str(record.get("doi") or "").strip() or None,
                    full_text_url=str(record.get("fullTextUrl") or "").strip() or None,
                    access_status=str(record.get("accessStatus") or "reviewed_full_text_excerpt").strip(),
                    warning=str(record.get("warning") or "").strip() or None,
                ))
        else:
            for evidence_id in reviewed_ids:
                row = by_id[evidence_id]
                evidence_records.append(EvidenceRecord(
                    provider=str(row.get("sourcePlatform") or "").strip(),
                    stable_id=evidence_id,
                    title=str(row.get("title") or "").strip(),
                    abstract=str(row.get("abstract") or "").strip(),
                    doi=str(row.get("doi") or "").strip() or None,
                    pmid=evidence_id.split(":", 1)[1] if evidence_id.startswith("pmid:") else None,
                    arxiv_id=evidence_id.split(":", 1)[1] if evidence_id.startswith("arxiv:") else None,
                    ads_id=evidence_id.split(":", 1)[1] if evidence_id.startswith("ads:") else None,
                    full_text_url=str(row.get("url") or "").strip() or None,
                    access_status=str(row.get("accessStatus") or "metadata").strip() or "metadata",
                    warning=str(row.get("warning") or "").strip(),
                ))
        selection = select_science125_evidence(question_id=question_id, query=query, records=tuple(evidence_records), max_selected=6)
        if reviewed_snapshot_records:
            families = tuple(sorted({
                str(record.get("providerFamily") or "").strip() or _provider_family(str(record.get("provider") or ""))
                for record in reviewed_snapshot_records if isinstance(record, dict)
            }))
            evidence_snapshot = {
                "selectionVersion": "science125-reviewed-interactive-v1",
                "evidenceStatus": "sufficient",
                "eligibleFullTextCount": len(reviewed_snapshot_records),
                "providerFamilyCount": len(families),
                "providerFamilies": list(families),
                "selectedEvidence": [
                    {
                        "stableId": record.get("stableId"),
                        "provider": record.get("provider"),
                        "providerFamily": record.get("providerFamily") or _provider_family(str(record.get("provider") or "")),
                        "title": record.get("title"),
                        "record": record,
                        "eligibleForGeneration": True,
                        "selected": True,
                        "reasons": ["USER_REVIEWED_FOR_SUCCESSFUL_INTERACTIVE_JOB"],
                    }
                    for record in reviewed_snapshot_records if isinstance(record, dict)
                ],
                "excludedEvidence": [],
            }
            selected_evidence_count = len(reviewed_snapshot_records)
            provider_families = families
        else:
            evidence_snapshot = selection.to_snapshot()
            selected_evidence_count = len(selection.selected)
            provider_families = selection.provider_families
        retrieval_snapshot = {
            "sourceType": "interactive_job",
            "searchJobId": literature_job.id,
            "query": query_info,
            "refinementQueries": search_result.get("refinementQueries") or [],
            "policyHashes": search_result.get("policyHashes") or {},
            "reviewedEvidenceIds": reviewed_ids,
            "sourceJobId": job.id,
        }
        selected = select_science125_hypothesis(output.model_dump(by_alias=True, mode="json"))
        provenance = {
            **(output.provenance.model_dump(by_alias=True, mode="json") if output.provenance else {}),
            **audit,
            "sourceType": "interactive_job",
            "sourceJobId": job.id,
        }
        prompt_version = str(audit.get("promptVersion") or SCIENCE125_PROMPT_VERSION)
        prompt_registry = load_science125_prompt_registry(prompt_version)
        with self._lock:
            existing = self._store.get_report_by_source_job(user_id=job.user_id, source_job_id=job.id)
            if existing is not None:
                self._ensure_report_exports(user_id=job.user_id, report=existing)
                return existing
            batch = self._store.create_batch(
                user_id=job.user_id,
                question_ids=(question_id,),
                manifest_version="science125-v1",
                manifest_sha256=get_science125_routing().base_manifest_content_sha256,
                routing_version="science125-routing-v1",
                routing_sha256=get_science125_routing().routing_content_sha256,
                prompt_version=prompt_version,
                prompt_registry_sha256=prompt_registry.registry_content_sha256,
                model=model,
            )
            try:
                report = self._store.upsert_report(
                    user_id=job.user_id,
                    batch_id=batch.id,
                    question_id=question_id,
                    question=item.question,
                    question_zh=item.question_zh,
                    benchmark_domain=item.benchmark_domain,
                    primary_subdomain=item.primary_subdomain,
                    attempt_number=1,
                    status="SUCCEEDED",
                    evidence_status="sufficient",
                    selected_evidence_count=selected_evidence_count,
                    provider_families=provider_families,
                    selected_hypothesis_id=selected.hypothesis_id,
                    selected_hypothesis_confidence=selected.confidence if selected.hypothesis_id else None,
                    selected_hypothesis_reason=selected.reason,
                    model=model,
                    request_id=str(audit.get("requestId") or output.provenance.request_id),
                    total_tokens=int(audit.get("totalTokens") or 0),
                    latency_ms=int(audit.get("latencyMs") or 0),
                    estimated_cost_cny=float(audit.get("estimatedCostCny") or 0.0),
                    retrieval_query=query,
                    retrieval_query_zh=None,
                    refinement_queries=tuple(str(value) for value in search_result.get("refinementQueries") or []),
                    retrieval_snapshot=retrieval_snapshot,
                    evidence_snapshot=evidence_snapshot,
                    evidence_snapshot_sha256=str(
                        private_snapshot.get("evidenceSnapshotHash")
                        or audit.get("evidenceSnapshotHash")
                        or _sha256_json(evidence_snapshot)
                    ),
                    research_output=output.model_dump(by_alias=True, mode="json"),
                    provenance=provenance,
                    source_type="interactive_job",
                    source_job_id=job.id,
                )
            except IntegrityError:
                existing = self._store.get_report_by_source_job(user_id=job.user_id, source_job_id=job.id)
                if existing is None:
                    raise
                report = existing
            self._store.update_batch_counters(user_id=job.user_id, batch_id=batch.id)
            self._ensure_report_exports(user_id=job.user_id, report=report)
            return report

    def _ensure_report_exports(self, *, user_id: int, report: StoredScience125Report) -> None:
        formats = {item.format for item in self._store.list_exports_for_report(user_id=user_id, report_id=report.id)}
        for format in ("json", "docx"):
            if format not in formats:
                self.create_report_export(user_id=user_id, report_id=report.id, format=format)

    def report_out(self, *, user_id: int, report_id: str) -> dict[str, Any] | None:
        report = self.get_report(user_id=user_id, report_id=report_id)
        if report is None:
            return None
        exports = tuple(item.to_dict() for item in self._store.list_exports_for_report(user_id=user_id, report_id=report_id))
        return _report_payload(report, exports)

    def list_reports_by_question(self, *, user_id: int, question_id: str) -> list[StoredScience125Report]:
        return self._store.list_reports_by_question(user_id=user_id, question_id=question_id)

    def pause_batch(self, *, user_id: int, batch_id: str) -> StoredScience125Batch | None:
        return self._store.set_batch_status(user_id=user_id, batch_id=batch_id, status="PAUSED")

    def resume_batch(self, *, user_id: int, batch_id: str) -> StoredScience125Batch | None:
        batch = self._store.mark_batch_running(user_id=user_id, batch_id=batch_id)
        if batch is None:
            return None
        self._submit_batch(batch_id=batch.id, user_id=user_id)
        return batch

    def retry_batch(
        self,
        *,
        user_id: int,
        batch_id: str,
        question_ids: tuple[str, ...] | None = None,
    ) -> StoredScience125Batch | None:
        existing_batch = self._store.get_batch_for_user(user_id=user_id, batch_id=batch_id)
        if existing_batch is None:
            return None
        retry_ids = self._store.prepare_items_for_retry(
            user_id=user_id,
            batch_id=batch_id,
            question_ids=question_ids,
        )
        if not retry_ids:
            return existing_batch
        batch = self._store.mark_batch_running(user_id=user_id, batch_id=batch_id)
        self._submit_batch(batch_id=batch.id, user_id=user_id, question_ids=retry_ids)
        return batch

    def _submit_batch(
        self,
        *,
        batch_id: str,
        user_id: int,
        question_ids: tuple[str, ...] | None = None,
    ) -> None:
        with self._lock:
            if batch_id in self._active_batch_ids:
                return
            self._active_batch_ids.add(batch_id)
        try:
            self._executor.submit(
                self._run_submitted_batch,
                batch_id=batch_id,
                user_id=user_id,
                question_ids=question_ids,
            )
        except RuntimeError:
            with self._lock:
                self._active_batch_ids.discard(batch_id)
            raise

    def _run_submitted_batch(
        self,
        *,
        batch_id: str,
        user_id: int,
        question_ids: tuple[str, ...] | None,
    ) -> StoredScience125Batch | None:
        try:
            return self._run_batch(batch_id=batch_id, user_id=user_id, question_ids=question_ids)
        finally:
            with self._lock:
                self._active_batch_ids.discard(batch_id)
            batch = self._store.get_batch_for_user(user_id=user_id, batch_id=batch_id)
            if batch is not None and batch.status == "RUNNING":
                items = self._store.list_batch_items(user_id=user_id, batch_id=batch_id)
                if any(item.status in {"PENDING", "RETRYING"} for item in items):
                    self._submit_batch(
                        batch_id=batch_id,
                        user_id=user_id,
                        question_ids=question_ids,
                    )

    def _required_llm_config(self, user_id: int):
        credential_environment = api_keys.api_key_store.science125_environment(user_id)
        settings = get_settings()
        api_key = str(credential_environment.get("DASHSCOPE_API_KEY") or "").strip()
        if not api_key:
            raise Science125ReportError(
                "DASHSCOPE_API_KEY_REQUIRED",
                "Science 125 generation requires a configured DashScope API Key.",
            )
        try:
            input_cost = float(str(
                credential_environment.get("QWEN_INPUT_COST_PER_MILLION_CNY")
                or settings.qwen_input_cost_per_million_cny
                or ""
            ))
            output_cost = float(str(
                credential_environment.get("QWEN_OUTPUT_COST_PER_MILLION_CNY")
                or settings.qwen_output_cost_per_million_cny
                or ""
            ))
        except ValueError as exc:
            raise Science125ReportError(
                "SCIENCE125_MODEL_PRICING_REQUIRED",
                "Science 125 generation requires valid Qwen input and output token prices.",
            ) from exc
        if input_cost <= 0 or output_cost <= 0:
            raise Science125ReportError(
                "SCIENCE125_MODEL_PRICING_REQUIRED",
                "Science 125 generation requires positive Qwen input and output token prices.",
            )
        return load_llm_client().LLMConfig(
            api_key=api_key,
            model=science125_model_name(),
            input_cost_per_million_cny=input_cost,
            output_cost_per_million_cny=output_cost,
        )

    def run_batch(
        self,
        *,
        batch_id: str,
        user_id: int | None = None,
        question_ids: tuple[str, ...] | None = None,
    ) -> StoredScience125Batch | None:
        with self._lock:
            if batch_id in self._active_batch_ids:
                raise Science125ReportError(
                    "SCIENCE125_BATCH_ALREADY_RUNNING",
                    "The Science 125 batch already has an active runner.",
                )
            self._active_batch_ids.add(batch_id)
        try:
            batch = self._store.get_batch(batch_id=batch_id)
            if batch is None:
                return None
            owner_id = user_id if user_id is not None else batch.user_id
            self._store.mark_batch_running(user_id=owner_id, batch_id=batch_id)
            return self._run_batch(batch_id=batch_id, user_id=user_id, question_ids=question_ids)
        finally:
            with self._lock:
                self._active_batch_ids.discard(batch_id)

    def _run_batch(
        self,
        *,
        batch_id: str,
        user_id: int | None = None,
        question_ids: tuple[str, ...] | None = None,
    ) -> StoredScience125Batch | None:
        batch = self._store.get_batch(batch_id=batch_id)
        if batch is None:
            return None
        owner_id = user_id if user_id is not None else batch.user_id
        if batch.status != "RUNNING":
            return batch
        items = self._store.list_batch_items_for_runner(batch_id=batch_id, question_ids=question_ids)
        consecutive_service_errors = 0
        for item in items:
            current = self._store.get_batch_for_user(user_id=owner_id, batch_id=batch_id)
            if current is None or current.status in {"PAUSED", "CANCELLED"}:
                break
            if item.status == "SUCCEEDED" and question_ids is None:
                continue
            try:
                self.run_report_item(user_id=owner_id, batch_id=batch_id, question_id=item.question_id, attempt_number=item.attempt_number)
                consecutive_service_errors = 0
            except Science125ReportError as exc:
                self._store.set_item_error(
                    user_id=owner_id,
                    batch_id=batch_id,
                    question_id=item.question_id,
                    status="FAILED",
                    code=exc.code,
                    message=exc.message,
                )
                is_service_error = exc.code in {
                    "DASHSCOPE_API_KEY_REQUIRED",
                    "SCIENCE125_MODEL_PRICING_REQUIRED",
                    "SCIENCE125_QWEN_FAILED",
                    "SCIENCE125_BATCH_ITEM_FAILED",
                }
                consecutive_service_errors = consecutive_service_errors + 1 if is_service_error else 0
                must_pause = (
                    exc.code in {"DASHSCOPE_API_KEY_REQUIRED", "SCIENCE125_MODEL_PRICING_REQUIRED"}
                    or "status=429" in exc.message
                    or "budget" in exc.message.casefold()
                    or consecutive_service_errors >= 3
                )
                if must_pause:
                    self._store.set_batch_status(user_id=owner_id, batch_id=batch_id, status="PAUSED")
                    break
            except Exception as exc:
                self._store.set_item_error(
                    user_id=owner_id,
                    batch_id=batch_id,
                    question_id=item.question_id,
                    status="FAILED",
                    code="SCIENCE125_BATCH_ITEM_FAILED",
                    message="Science 125 batch item failed unexpectedly.",
                )
                consecutive_service_errors += 1
                if consecutive_service_errors >= 3:
                    self._store.set_batch_status(user_id=owner_id, batch_id=batch_id, status="PAUSED")
                    break
        return self._store.update_batch_counters(user_id=owner_id, batch_id=batch_id)

    def run_report_item(
        self,
        *,
        user_id: int,
        batch_id: str,
        question_id: str,
        attempt_number: int = 1,
    ) -> StoredScience125Report:
        batch = self._store.get_batch_for_user(user_id=user_id, batch_id=batch_id)
        if batch is None:
            raise Science125ReportError("SCIENCE125_BATCH_NOT_FOUND", "Science 125 batch not found.")
        try:
            pinned_registry = load_science125_prompt_registry(batch.prompt_version)
        except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
            raise Science125ReportError(
                "SCIENCE125_PROMPT_SNAPSHOT_UNAVAILABLE",
                "The prompt registry pinned to this Science 125 batch is unavailable.",
            ) from exc
        if pinned_registry.registry_content_sha256 != batch.prompt_registry_sha256:
            raise Science125ReportError(
                "SCIENCE125_PROMPT_SNAPSHOT_MISMATCH",
                "The prompt registry no longer matches the hash pinned to this Science 125 batch.",
            )
        catalog = get_science125_catalog()
        catalog_item = next((item for item in catalog.data if item.id == question_id), None)
        if catalog_item is None:
            raise Science125ReportError("SCIENCE125_QUESTION_NOT_FOUND", "Science 125 question not found.")
        context_index = load_science125_context_index()
        context_item = context_index.items.get(question_id)
        if context_item is None:
            raise Science125ReportError("SCIENCE125_CONTEXT_UNAVAILABLE", "Science 125 source context is unavailable.")
        route = get_science125_route(question_id)
        retrieval_profile = get_science125_retrieval_profile(route.retrieval_profile)
        localization = get_science125_localization(question_id)
        query = (
            str(localization.recommended_query).strip()
            if localization and localization.recommended_query
            else build_science125_search_query(catalog_item.question, context_item.source_context)
        )
        query_zh = localization.search_intent_zh if localization else None
        self._store.upsert_report(
            user_id=user_id,
            batch_id=batch_id,
            question_id=question_id,
            question=catalog_item.question,
            question_zh=catalog_item.question_zh,
            benchmark_domain=catalog_item.benchmark_domain,
            primary_subdomain=catalog_item.primary_subdomain,
            attempt_number=attempt_number,
            status="RETRIEVING",
            evidence_status="insufficient",
            selected_evidence_count=0,
            retrieval_query=query,
            retrieval_query_zh=query_zh,
        )

        credential_environment = api_keys.api_key_store.science125_environment(user_id)
        adapters = default_provider_adapters(environ=credential_environment)
        refinement_queries = build_science125_refinement_queries(
            retrieval_profile.query_adapter,
            query,
            primary_subdomain=route.primary_subdomain,
        )
        queries = (query, *refinement_queries)
        all_records: list[EvidenceRecord] = []
        diagnostics: list[dict[str, Any]] = []
        query_results: list[dict[str, Any]] = []
        selection: Science125EvidenceSelection | None = None
        for current_query in queries:
            result = search_science125(
                route.retrieval_profile,
                current_query,
                adapters=adapters,
                store=self._rate_store,
                environ=credential_environment,
                user_id=user_id,
            )
            records = list(result.evidence)
            all_records.extend(records)
            diagnostics.extend(diagnostic.to_dict() for diagnostic in result.diagnostics)
            query_results.append({
                "query": current_query,
                "queryHash": result.query_hash,
                "cacheKey": result.cache_key,
                "resultCount": len(records),
                "diagnostics": [diagnostic.to_dict() for diagnostic in result.diagnostics],
            })
            scoring_query = " ".join(str(item["query"]) for item in query_results)
            selection = select_science125_evidence(
                question_id=question_id,
                query=scoring_query,
                records=tuple(all_records),
                max_selected=6,
            )
            if selection.evidence_status == "sufficient":
                break
        if selection is None:
            selection = select_science125_evidence(question_id=question_id, query=query, records=(), max_selected=6)
        retrieval_snapshot = {
            "profileId": route.retrieval_profile,
            "queryAdapter": retrieval_profile.query_adapter,
            "queries": query_results,
            "diagnostics": diagnostics,
            "resultCount": len(all_records),
        }
        evidence_snapshot = selection.to_snapshot()
        if selection.evidence_status != "sufficient":
            report = self._store.upsert_report(
                user_id=user_id,
                batch_id=batch_id,
                question_id=question_id,
                question=catalog_item.question,
                question_zh=catalog_item.question_zh,
                benchmark_domain=catalog_item.benchmark_domain,
                primary_subdomain=catalog_item.primary_subdomain,
                attempt_number=attempt_number,
                status="BLOCKED_EVIDENCE",
                evidence_status="insufficient",
                selected_evidence_count=0,
                provider_families=selection.provider_families,
                retrieval_query=query,
                retrieval_query_zh=query_zh,
                refinement_queries=refinement_queries,
                retrieval_snapshot=retrieval_snapshot,
                evidence_snapshot=evidence_snapshot,
                evidence_snapshot_sha256=selection.evidence_snapshot_hash,
            )
            self._store.update_batch_counters(user_id=user_id, batch_id=batch_id)
            return report

        report = self._store.upsert_report(
            user_id=user_id,
            batch_id=batch_id,
            question_id=question_id,
            question=catalog_item.question,
            question_zh=catalog_item.question_zh,
            benchmark_domain=catalog_item.benchmark_domain,
            primary_subdomain=catalog_item.primary_subdomain,
            attempt_number=attempt_number,
            status="GENERATING",
            evidence_status="sufficient",
            selected_evidence_count=len(selection.selected),
            provider_families=selection.provider_families,
            retrieval_query=query,
            retrieval_query_zh=query_zh,
            refinement_queries=refinement_queries,
            retrieval_snapshot=retrieval_snapshot,
            evidence_snapshot=evidence_snapshot,
            evidence_snapshot_sha256=selection.evidence_snapshot_hash,
        )
        config = self._required_llm_config(user_id)
        llm = load_llm_client()
        budget = llm.LLMBudget(max_total_tokens=20_000, max_estimated_cost_cny=3.0)
        request = ResearchGenerationRequest(
            question=catalog_item.question,
            profile="general_science",
            chemistry_subdomain=None,
            candidate_count=3,
            science125_id=question_id,
            science125_source_context=context_item.source_context,
            science125_routing={
                **route.model_dump(by_alias=True),
                "routingVersion": "science125-routing-v1",
            },
            science125_prompt_version=batch.prompt_version,
            evidence_records=selection.evidence_records(),
        )

        def telemetry_sink(event: Mapping[str, Any]) -> None:
            enriched = dict(event)
            enriched["policy_hash"] = _sha256_json(retrieval_snapshot)
            enriched["evidence_snapshot_hash"] = selection.evidence_snapshot_hash
            self._ledger.record(enriched)

        try:
            generated = ResearchGenerationService.from_environment().generate(
                request,
                config=config,
                budget=budget,
                context=llm.LLMCallContext(resource_type="science125_report", resource_id=report.id),
                telemetry_sink=telemetry_sink,
            )
        except ResearchGenerationValidationError as exc:
            raise Science125ReportError(
                "SCIENCE125_SCHEMA_INVALID",
                "Qwen returned a Science 125 result that could not satisfy the research-v1 contract.",
            ) from exc
        except ResearchGenerationCallError as exc:
            result = exc.result
            raise Science125ReportError(
                "SCIENCE125_QWEN_FAILED",
                f"DashScope Qwen generation failed: status={result.status_code or 'none'}, "
                f"error_type={result.error_type or result.status}, attempts={result.attempts}.",
            ) from exc
        output = generated.output.model_dump(by_alias=True, mode="json")
        selected = select_science125_hypothesis(output)
        final_report = self._store.upsert_report(
            user_id=user_id,
            batch_id=batch_id,
            question_id=question_id,
            question=catalog_item.question,
            question_zh=catalog_item.question_zh,
            benchmark_domain=catalog_item.benchmark_domain,
            primary_subdomain=catalog_item.primary_subdomain,
            attempt_number=attempt_number,
            status="SUCCEEDED",
            evidence_status="sufficient",
            selected_evidence_count=len(selection.selected),
            provider_families=selection.provider_families,
            selected_hypothesis_id=selected.hypothesis_id,
            selected_hypothesis_confidence=selected.confidence if selected.hypothesis_id else None,
            selected_hypothesis_reason=selected.reason,
            model=generated.call.model,
            request_id=generated.call.request_id,
            total_tokens=generated.total_tokens,
            latency_ms=generated.latency_ms,
            estimated_cost_cny=generated.estimated_cost_cny,
            retrieval_query=query,
            retrieval_query_zh=query_zh,
            refinement_queries=refinement_queries,
            retrieval_snapshot=retrieval_snapshot,
            evidence_snapshot=evidence_snapshot,
            evidence_snapshot_sha256=selection.evidence_snapshot_hash,
            research_output=output,
            provenance={
                "provider": generated.call.provider,
                "model": generated.call.model,
                "requestId": generated.call.request_id,
                "totalTokens": generated.total_tokens,
                "latencyMs": generated.latency_ms,
                "retryCount": generated.retry_count,
                "estimatedCostCny": generated.estimated_cost_cny,
                "schemaRepaired": generated.schema_repaired,
            },
        )
        self.create_report_export(user_id=user_id, report_id=final_report.id, format="json")
        self._store.update_batch_counters(user_id=user_id, batch_id=batch_id)
        return final_report

    def create_report_export(self, *, user_id: int, report_id: str, format: str) -> StoredScience125Export:
        report = self._store.get_report_for_user(user_id=user_id, report_id=report_id)
        if report is None:
            raise Science125ReportError("NOT_FOUND", "Science 125 report not found.")
        if format == "docx" and report.status != "SUCCEEDED":
            raise Science125ReportError("REPORT_NOT_SUCCEEDED", "Only succeeded reports can be exported as DOCX.")
        root = science125_exports_dir() / str(user_id) / report.batch_id
        suffix = "docx" if format == "docx" else "json"
        path = root / f"{report.question_id}-{report.id[:8]}.{suffix}"
        if format == "docx":
            _build_report_docx(path, report)
            mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif format == "json":
            _write_json(path, _report_payload(report))
            mime = "application/json"
        else:
            raise Science125ReportError("UNSUPPORTED_EXPORT_FORMAT", "Unsupported Science 125 export format.")
        return self._store.create_export(
            user_id=user_id,
            batch_id=report.batch_id,
            report_id=report.id,
            format=format,
            file_name=path.name,
            mime_type=mime,
            size_bytes=path.stat().st_size,
            managed_path=path,
        )

    def create_batch_export(self, *, user_id: int, batch_id: str, format: str) -> StoredScience125Export:
        batch = self._store.get_batch_for_user(user_id=user_id, batch_id=batch_id)
        if batch is None:
            raise Science125ReportError("NOT_FOUND", "Science 125 batch not found.")
        reports = self._store.list_reports(user_id=user_id, batch_id=batch_id, sort_by="questionId", sort_order="asc")
        if format == "docx":
            if batch.total_count != 125 or batch.succeeded_count != 125 or len(reports) != 125:
                raise Science125ReportError(
                    "BATCH_FINAL_DOCX_NOT_READY",
                    "The final Science 125 DOCX requires 125/125 succeeded reports.",
                )
            path = science125_exports_dir() / str(user_id) / batch.id / f"science125-final-{batch.id[:8]}.docx"
            _build_batch_docx(path, batch, reports)
            mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif format == "json":
            path = science125_exports_dir() / str(user_id) / batch.id / f"science125-audit-{batch.id[:8]}.json"
            _write_json(path, _batch_payload(batch, tuple(reports)))
            mime = "application/json"
        else:
            raise Science125ReportError("UNSUPPORTED_EXPORT_FORMAT", "Unsupported Science 125 export format.")
        return self._store.create_export(
            user_id=user_id,
            batch_id=batch.id,
            report_id=None,
            format=format,
            file_name=path.name,
            mime_type=mime,
            size_bytes=path.stat().st_size,
            managed_path=path,
        )

    def get_export(self, *, user_id: int, export_id: str) -> StoredScience125Export | None:
        export = self._store.get_export_for_user(user_id=user_id, export_id=export_id)
        if export is None:
            return None
        root = (science125_exports_dir() / str(user_id)).resolve()
        if export.path.is_symlink() or not export.path.resolve().is_relative_to(root):
            raise Science125ReportError("UNSAFE_EXPORT_PATH", "Science 125 export path is outside managed storage.")
        return export


@lru_cache(maxsize=1)
def get_science125_report_service() -> Science125ReportService:
    return Science125ReportService()


__all__ = [
    "SCIENCE125_REPORT_EXPORT_DIR",
    "Science125EvidenceSelection",
    "Science125EvidenceSelectionItem",
    "Science125SelectedHypothesis",
    "science125_exports_dir",
    "science125_model_name",
    "Science125ReportError",
    "Science125ReportService",
    "select_science125_evidence",
    "select_science125_hypothesis",
    "get_science125_report_service",
    "build_science125_search_query",
]
